import os
from io import BytesIO
from pathlib import Path

from bs4 import BeautifulSoup
from docx import Document
from PyPDF2 import PdfReader
from pptx import Presentation


TEXT_EXTENSIONS = {
    ".txt",
    ".md",
    ".csv",
    ".json",
    ".xml",
    ".html",
    ".htm",
    ".log",
    ".py",
    ".js",
    ".ts",
    ".css",
    ".java",
    ".c",
    ".cpp",
    ".rb",
    ".go",
    ".php",
    ".sql",
    ".yaml",
    ".yml",
    ".ini",
    ".rtf",
}
MAX_UPLOAD_BYTES = int(os.getenv("MAX_UPLOAD_BYTES", 5 * 1024 * 1024))


def _read_text_bytes(data: bytes) -> str:
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("Unable to decode the uploaded file as text.")


def _extract_html_text(data: bytes) -> str:
    html = _read_text_bytes(data)
    soup = BeautifulSoup(html, "html.parser")
    return soup.get_text(separator=" ", strip=True)


def _extract_docx_text(data: bytes) -> str:
    document = Document(BytesIO(data))
    return "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text.strip())


def _extract_pdf_text(data: bytes) -> str:
    reader = PdfReader(BytesIO(data))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _extract_pptx_text(data: bytes) -> str:
    presentation = Presentation(BytesIO(data))
    chunks = []

    for slide in presentation.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                chunks.append(shape.text)

    return "\n".join(chunk for chunk in chunks if chunk.strip())


def extract_text_from_upload(uploaded_file) -> tuple[str, str]:
    filename = uploaded_file.filename or "uploaded-file"
    suffix = Path(filename).suffix.lower()
    data = uploaded_file.read()

    if not data:
        raise ValueError("The uploaded file is empty.")
    if len(data) > MAX_UPLOAD_BYTES:
        raise ValueError("The uploaded file is too large. The maximum size is 5 MB.")

    if suffix == ".ppt":
        raise ValueError("Legacy .ppt files are not supported yet. Please convert the presentation to .pptx and upload it again.")
    try:
        if suffix == ".pdf":
            text = _extract_pdf_text(data)
        elif suffix == ".docx":
            text = _extract_docx_text(data)
        elif suffix == ".pptx":
            text = _extract_pptx_text(data)
        elif suffix in {".html", ".htm", ".xml"}:
            text = _extract_html_text(data)
        elif suffix in TEXT_EXTENSIONS or not suffix:
            text = _read_text_bytes(data)
        else:
            raise ValueError("Unsupported file type. Upload TXT, PDF, DOCX, PPTX, HTML, CSV, JSON, code, or log files.")
    except ValueError:
        raise
    except Exception as error:
        raise ValueError("The file could not be read. Check that it is a valid, non-encrypted document.") from error

    normalized = " ".join(text.split())
    if not normalized:
        raise ValueError("No readable text was found in the uploaded file.")

    return normalized, filename
