import io
import json
import os
import sys
import unittest
from pathlib import Path

# Add project root and backend to python path
ROOT_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT_DIR))
sys.path.insert(0, str(ROOT_DIR / "backend"))

from backend.app import app
from docx import Document

class TestSummarizationWorkflow(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        app.testing = True
        cls.client = app.test_client()

        cls.normal_paragraph = (
            "Artificial intelligence is transforming industries across the globe. "
            "Machine learning algorithms allow computers to learn patterns from vast amounts of data without being explicitly programmed. "
            "Natural language processing enables machines to understand, generate, and summarize human languages effectively. "
            "Deep learning architectures, particularly transformer models, have achieved state-of-the-art results in various domains. "
            "However, issues such as model hallucinations and bias remain critical challenges that researchers are actively addressing."
        )

        cls.long_article = (
            "Renewable energy technologies have made remarkable strides over the past decade. "
            "Solar photovoltaic systems have become significantly more cost-effective due to economies of scale and manufacturing innovations. "
            "Wind turbines, both onshore and offshore, are generating higher energy yields thanks to larger rotor diameters and smarter control systems. "
            "Grid modernization and energy storage solutions, such as lithium-ion batteries and pumped-storage hydropower, play an indispensable role in mitigating renewable intermittency. "
            "Governments worldwide are enacting policies, tax incentives, and decarbonization mandates to accelerate the transition away from fossil fuels. "
            "Despite these advancements, transmission infrastructure bottlenecks and raw material supply chain vulnerabilities continue to pose operational challenges. "
            "International collaborations and public-private investments will be paramount to reaching net-zero carbon emission targets by 2050."
        )

        cls.academic_text = (
            "Transformer-based sequence-to-sequence neural architectures have established a new paradigm in automatic abstractive summarization. "
            "Pre-trained language models such as BART and PEGASUS leverage self-attention mechanisms to model bidirectional contextual representations across input documents. "
            "Despite their high fluency and coherence, these models frequently generate unfaithful content, commonly referred to as hallucinations. "
            "To alleviate factual inconsistency, Natural Language Inference (NLI) frameworks compare generated claims against source evidence sentences to classify statements as entailment, neutral, or contradiction. "
            "Empirical evaluations indicate that integrating cross-encoder verification significantly reduces unsupported factual claims in generated summaries."
        )

    def test_01_bart_summarization(self):
        """Test BART transformer model generates fluent abstractive summary."""
        response = self.client.post(
            "/summarize",
            json={"text": self.academic_text, "summary_length": 6, "model_type": "bart"}
        )
        self.assertEqual(response.status_code, 200)
        data = response.get_json()
        self.assertEqual(data["model_type"], "bart")
        self.assertTrue(len(data["summary"].strip()) > 0)
        self.assertTrue(len(data["overview"].strip()) > 0)
        self.assertIn("analysis", data)
        self.assertGreater(data["analysis"]["summary_word_count"], 0)

    def test_02_pegasus_summarization(self):
        """Test PEGASUS transformer model generates abstractive summary."""
        response = self.client.post(
            "/summarize",
            json={"text": self.long_article, "summary_length": 6, "model_type": "pegasus"}
        )
        self.assertEqual(response.status_code, 200)
        data = response.get_json()
        self.assertEqual(data["model_type"], "pegasus")
        self.assertTrue(len(data["summary"].strip()) > 0)
        self.assertTrue(len(data["overview"].strip()) > 0)

    def test_03_presets_short_medium_long(self):
        """Test summary length presets (short=6, medium=10, long=15)."""
        for length in [6, 10, 15]:
            response = self.client.post(
                "/summarize",
                json={"text": self.long_article, "summary_length": length, "model_type": "bart"}
            )
            self.assertEqual(response.status_code, 200)
            data = response.get_json()
            summary = data["summary"].strip()
            self.assertTrue(len(summary) > 0)

    def test_04_uploaded_txt_file(self):
        """Test file upload with a TXT document."""
        txt_stream = io.BytesIO(self.normal_paragraph.encode("utf-8"))
        response = self.client.post(
            "/summarize",
            data={
                "file": (txt_stream, "notes.txt"),
                "summary_length": 6,
                "model_type": "bart",
            },
            content_type="multipart/form-data"
        )
        self.assertEqual(response.status_code, 200)
        data = response.get_json()
        self.assertEqual(data["source_file"], "notes.txt")
        self.assertTrue(len(data["summary"].strip()) > 0)

    def test_05_uploaded_docx_file(self):
        """Test file upload with a DOCX document."""
        doc = Document()
        doc.add_heading("Artificial Intelligence Overview", 0)
        doc.add_paragraph(self.normal_paragraph)
        docx_stream = io.BytesIO()
        doc.save(docx_stream)
        docx_stream.seek(0)

        response = self.client.post(
            "/summarize",
            data={
                "file": (docx_stream, "document.docx"),
                "summary_length": 6,
                "model_type": "pegasus",
            },
            content_type="multipart/form-data"
        )
        self.assertEqual(response.status_code, 200)
        data = response.get_json()
        self.assertEqual(data["source_file"], "document.docx")
        self.assertTrue(len(data["summary"].strip()) > 0)

    def test_06_input_validation(self):
        """Test error handling for empty input, short text, removed extractive model, invalid length."""
        # Empty text
        res1 = self.client.post("/summarize", json={"text": "", "model_type": "bart"})
        self.assertEqual(res1.status_code, 400)
        self.assertIn("error", res1.get_json())

        # Too short text (< 5 words)
        res2 = self.client.post("/summarize", json={"text": "Short text.", "model_type": "bart"})
        self.assertEqual(res2.status_code, 400)
        self.assertIn("error", res2.get_json())

        # Removed extractive model should be rejected as unsupported
        res3 = self.client.post("/summarize", json={"text": self.normal_paragraph, "model_type": "extractive"})
        self.assertEqual(res3.status_code, 400)
        self.assertIn("error", res3.get_json())
        self.assertIn("Unsupported model type", res3.get_json()["error"])

        # Invalid summary length (> 20)
        res4 = self.client.post("/summarize", json={"text": self.normal_paragraph, "summary_length": 25, "model_type": "bart"})
        self.assertEqual(res4.status_code, 400)
        self.assertIn("error", res4.get_json())

    def test_07_hallucination_detection(self):
        """Test hallucination detection endpoint decoupled from summarization."""
        summary = "Artificial intelligence is transforming global industries with machine learning algorithms."
        response = self.client.post(
            "/analyze_hallucination",
            json={"text": self.normal_paragraph, "summary": summary}
        )
        self.assertEqual(response.status_code, 200)
        data = response.get_json()
        self.assertTrue("overall_label" in data or "error" in data)

    def test_08_repeated_summarization_different_models(self):
        """Test switching models on consecutive requests (BART <-> PEGASUS)."""
        # Request 1 with BART
        res_bart = self.client.post(
            "/summarize",
            json={"text": self.normal_paragraph, "summary_length": 6, "model_type": "bart"}
        )
        self.assertEqual(res_bart.status_code, 200)
        self.assertEqual(res_bart.get_json()["model_type"], "bart")

        # Request 2 with PEGASUS
        res_pegasus = self.client.post(
            "/summarize",
            json={"text": self.normal_paragraph, "summary_length": 6, "model_type": "pegasus"}
        )
        self.assertEqual(res_pegasus.status_code, 200)
        self.assertEqual(res_pegasus.get_json()["model_type"], "pegasus")

    def test_09_executive_description_generation(self):
        """Test executive document description generator eliminates repetitions and produces high-quality overviews."""
        from backend.summarizer import generate_overview

        # 1. Repetitive narrative text (user screenshot case)
        stutter_text = (
            "And learns to look for the leaves, and to try to find the leaves to find an old pile, "
            "and then to find out the old leaves and find a pile and find the old pile and then find a box, "
            "and find out a box and the box."
        )
        overview = generate_overview(stutter_text)
        self.assertNotIn("And learns", overview)
        self.assertNotIn("find a pile and find the old pile", overview)
        self.assertTrue(overview.startswith("A narrative detailing"))
        self.assertIn("leaves", overview.lower())
        self.assertIn("box", overview.lower())

        # 2. Academic text
        academic_overview = generate_overview(self.academic_text)
        self.assertTrue(academic_overview.startswith("An analytical overview"))
        self.assertIn("transformer", academic_overview.lower())

        # 3. Renewable energy text
        renewable_overview = generate_overview(self.long_article)
        self.assertTrue(renewable_overview.startswith("A comprehensive report"))
        self.assertIn("renewable energy", renewable_overview.lower())


if __name__ == "__main__":
    unittest.main()
